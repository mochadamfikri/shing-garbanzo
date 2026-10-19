"""Endpoints for Daily, Weekly, Monthly Reports and genuine .xlsx Excel Export."""
import io
from datetime import datetime, timezone, timedelta
from typing import Literal
from fastapi import APIRouter, Depends, Query, Response, HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from auth import require_roles, ROLE_OWNER, ROLE_ADMIN
from utils import now_utc, format_tanggal_id, format_rupiah

router = APIRouter(prefix="/api/reports-v3", tags=["reports-v3"])

PeriodType = Literal["daily", "weekly", "monthly"]


def get_period_bounds(period: PeriodType) -> tuple[datetime, datetime, str]:
    now = now_utc()
    # End of current day in UTC
    end = now.replace(hour=23, minute=59, second=59, microsecond=999999)
    if period == "daily":
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)
        label = f"Harian ({start.strftime('%d %b %Y')})"
    elif period == "weekly":
        # Last 7 days
        start = (now - timedelta(days=6)).replace(hour=0, minute=0, second=0, microsecond=0)
        label = f"Mingguan ({start.strftime('%d %b')} - {end.strftime('%d %b %Y')})"
    elif period == "monthly":
        # Current month (from 1st of month)
        start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        label = f"Bulanan ({start.strftime('%B %Y')})"
    return start, end, label


async def fetch_report_data(db, start: datetime, end: datetime):
    # Fetch POS sales in range
    q_sales = {
        "status": "SELESAI",
        "tanggal_jual": {"$gte": start.isoformat(), "$lte": end.isoformat()}
    }
    sales = await db.pos_sales.find(q_sales).sort("tanggal_jual", -1).to_list(2000)

    # Fetch Expenses in range
    q_exp = {
        "tanggal": {"$gte": start.strftime("%Y-%m-%d"), "$lte": end.strftime("%Y-%m-%d")}
    }
    expenses = await db.expenses.find(q_exp).to_list(1000)

    total_omzet = sum(float(s.get("penjualan_bersih") or 0) for s in sales)
    total_hpp = sum(float(s.get("hpp") or 0) for s in sales)
    total_laba_kotor = sum(float(s.get("laba_kotor") or (float(s.get("penjualan_bersih") or 0) - float(s.get("hpp") or 0))) for s in sales)
    total_pengeluaran = sum(float(e.get("nominal") or 0) for e in expenses)
    laba_bersih = total_laba_kotor - total_pengeluaran

    # Total units sold & detailed list of units
    sold_units = []
    for s in sales:
        for item in s.get("items", []):
            sold_units.append({
                "nomor_transaksi": s.get("nomor_transaksi"),
                "tanggal": s.get("tanggal_jual"),
                "merek": item.get("merek", ""),
                "model": item.get("model", ""),
                "imei": item.get("imei_masked") or item.get("imei_1", "-"),
                "harga_jual": float(item.get("harga_akhir") or item.get("harga_jual") or 0),
                "modal": float(item.get("modal") or 0),
                "laba": float((item.get("harga_akhir") or item.get("harga_jual") or 0) - (item.get("modal") or 0)),
                "kasir": s.get("kasir", "Kasir"),
                "metode": s.get("metode_pembayaran", "TUNAI")
            })

    return {
        "sales": sales,
        "expenses": expenses,
        "sold_units": sold_units,
        "summary": {
            "total_transaksi": len(sales),
            "total_unit_terjual": len(sold_units),
            "total_omzet": total_omzet,
            "total_hpp": total_hpp,
            "total_laba_kotor": total_laba_kotor,
            "total_pengeluaran": total_pengeluaran,
            "laba_bersih": laba_bersih,
        }
    }


@router.get("/period")
async def get_period_report(
    period: PeriodType = Query("daily"),
    _user=Depends(require_roles(ROLE_OWNER, ROLE_ADMIN)),
):
    from server import db
    start, end, label = get_period_bounds(period)
    data = await fetch_report_data(db, start, end)
    return {
        "period": period,
        "period_label": label,
        "start": start.isoformat(),
        "end": end.isoformat(),
        "summary": data["summary"],
        "sold_units": data["sold_units"][:100],  # Limit preview items
        "expenses": data["expenses"][:50],
        "sales": [
            {
                "id": s.get("id"),
                "nomor_transaksi": s.get("nomor_transaksi"),
                "penjualan_bersih": s.get("penjualan_bersih"),
                "metode_pembayaran": s.get("metode_pembayaran"),
                "tanggal_jual": s.get("tanggal_jual"),
                "items_count": len(s.get("items", []))
            }
            for s in data["sales"][:50]
        ]
    }


@router.get("/export-excel")
async def export_excel(
    period: PeriodType = Query("daily"),
    _user=Depends(require_roles(ROLE_OWNER, ROLE_ADMIN)),
):
    from server import db
    start, end, label = get_period_bounds(period)
    data = await fetch_report_data(db, start, end)

    wb = Workbook()

    # Styling definitions
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    font_bold = Font(name="Calibri", size=11, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # 1. Sheet: Ringkasan
    ws_summary = wb.active
    ws_summary.title = "Ringkasan"
    ws_summary.append(["LAPORAN REKAP KEUANGAN & PENJUALAN - SADAM GADGET"])
    ws_summary.append([f"Periode: {label}"])
    ws_summary.append([f"Diekspor pada: {now_utc().strftime('%d/%m/%Y %H:%M:%S UTC')}"])
    ws_summary.append([])

    headers_summary = ["Indikator / Metrik", "Nilai (Rp / Jumlah)"]
    ws_summary.append(headers_summary)
    for col in range(1, 3):
        c = ws_summary.cell(row=5, column=col)
        c.font = font_header
        c.fill = fill_header

    metrics = [
        ("Total Transaksi POS", data["summary"]["total_transaksi"]),
        ("Total Unit HP Terjual", data["summary"]["total_unit_terjual"]),
        ("Total Omzet Penjualan", data["summary"]["total_omzet"]),
        ("Total Modal (HPP)", data["summary"]["total_hpp"]),
        ("Laba Kotor (Gross Profit)", data["summary"]["total_laba_kotor"]),
        ("Total Biaya Pengeluaran Operasional", data["summary"]["total_pengeluaran"]),
        ("Laba Bersih (Net Profit)", data["summary"]["laba_bersih"]),
    ]
    for row_idx, (m, v) in enumerate(metrics, start=6):
        ws_summary.append([m, v])
        ws_summary.cell(row=row_idx, column=1).border = border_thin
        cell_v = ws_summary.cell(row=row_idx, column=2)
        cell_v.border = border_thin
        if "Total Transaksi" not in m and "Total Unit" not in m:
            cell_v.number_format = '#,##0'

    ws_summary.column_dimensions['A'].width = 38
    ws_summary.column_dimensions['B'].width = 24

    # 2. Sheet: Transaksi
    ws_tx = wb.create_sheet(title="Transaksi")
    headers_tx = ["No. Transaksi", "Waktu Transaksi", "Kasir", "Pelanggan", "Metode Bayar", "Jumlah Unit", "Total Penjualan", "HPP Modal", "Laba Kotor"]
    ws_tx.append(headers_tx)
    for col in range(1, len(headers_tx) + 1):
        c = ws_tx.cell(row=1, column=col)
        c.font = font_header
        c.fill = fill_header

    for s in data["sales"]:
        ws_tx.append([
            s.get("nomor_transaksi"),
            s.get("tanggal_jual", "")[:19].replace("T", " "),
            s.get("kasir", "-"),
            s.get("pelanggan", "-"),
            s.get("metode_pembayaran"),
            len(s.get("items", [])),
            float(s.get("penjualan_bersih") or 0),
            float(s.get("hpp") or 0),
            float(s.get("laba_kotor") or 0)
        ])
    for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
        ws_tx.column_dimensions[col_letter].width = 18

    # 3. Sheet: Unit Terjual
    ws_units = wb.create_sheet(title="Unit Terjual")
    headers_units = ["No. Transaksi", "Merek", "Model Unit", "IMEI", "Harga Jual", "Modal Beli", "Laba Unit", "Metode", "Kasir"]
    ws_units.append(headers_units)
    for col in range(1, len(headers_units) + 1):
        c = ws_units.cell(row=1, column=col)
        c.font = font_header
        c.fill = fill_header

    for u in data["sold_units"]:
        ws_units.append([
            u["nomor_transaksi"],
            u["merek"],
            u["model"],
            u["imei"],
            u["harga_jual"],
            u["modal"],
            u["laba"],
            u["metode"],
            u["kasir"]
        ])
    for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
        ws_units.column_dimensions[col_letter].width = 18

    # 4. Sheet: Rekap Keuangan
    ws_fin = wb.create_sheet(title="Rekap Keuangan")
    ws_fin.append(["Komponen", "Nominal (Rp)"])
    for col in range(1, 3):
        c = ws_fin.cell(row=1, column=col)
        c.font = font_header
        c.fill = fill_header

    ws_fin.append(["Omzet Penjualan Kasir", data["summary"]["total_omzet"]])
    ws_fin.append(["(-) Beban Pokok Penjualan (HPP)", data["summary"]["total_hpp"]])
    ws_fin.append(["(=) Laba Kotor", data["summary"]["total_laba_kotor"]])
    ws_fin.append(["(-) Pengeluaran Operasional Toko", data["summary"]["total_pengeluaran"]])
    ws_fin.append(["(=) Laba Bersih Akhir", data["summary"]["laba_bersih"]])
    for r in range(2, 7):
        ws_fin.cell(row=r, column=2).number_format = '#,##0'
    ws_fin.column_dimensions['A'].width = 35
    ws_fin.column_dimensions['B'].width = 24

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Laporan_Sadam_Gadget_{period}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
